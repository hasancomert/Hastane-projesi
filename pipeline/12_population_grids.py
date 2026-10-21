"""Faz 1 — nüfus ızgaraları: GHS-POP R2023A (100 m ve 1 km), WorldPop R2025A (100 m), Kontur Population (H3 r8).

Adımlar
  1. Resmî JRC karoları (R5_C21) indirilir; Ankara penceresi (common.GHS_WINDOW) kırpılır, nodata (-200) -> 0.
     JRC'nin AWS kopyası (jrc-ghsl.s3) aynı pencerede piksel piksel karşılaştırılır.
  2. WorldPop R2025A constrained 100 m (2025, 2030) yerel WGS84 3" ızgarasında Ankara kutusuna kırpılır.
  3. Kontur Population TR 2023-11-01 (H3 çözünürlük 8) okunur.
  4. İlçe toplamları (piksel merkezi poligonda kuralı) -> data/processed/ilce_nufus_kaynaklar.csv.
     TÜİK değerleri burada YOK; entegrasyon adımında birleştirilir.

Çıktılar: data/processed/pop/*.tif, data/processed/ilce_nufus_kaynaklar.csv, data/raw/nufus/nufus_izgara_ozet.json
Önkoşul: pipeline/11_boundaries.py (data/processed/ankara_ilceler.geojson, ankara_il.geojson).

Çalıştırma (repo kökünden):  .venv/bin/python pipeline/12_population_grids.py
"""
from __future__ import annotations

import gzip
import hashlib
import json
import sys
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pipeline import common  # noqa: E402

RAW_N = common.RAW / "nufus"
GHS_DIR = RAW_N / "ghs"
WP_DIR = RAW_N / "worldpop"
KONTUR_DIR = RAW_N / "kontur"
POP_DIR = common.PROCESSED / "pop"
OUT_CSV = common.PROCESSED / "ilce_nufus_kaynaklar.csv"
SUMMARY_JSON = RAW_N / "nufus_izgara_ozet.json"
S3_CHECK_JSON = GHS_DIR / "s3_kopya_kontrol.json"

GHS_BASE = "https://jeodpp.jrc.ec.europa.eu/ftp/jrc-opendata/GHSL/GHS_POP_GLOBE_R2023A"
GHS_TILE = "R5_C21"
GHS_RES = {100: "100", 1000: "1000"}


def ghs_name(epoch: int, res: int) -> str:
    return f"GHS_POP_E{epoch}_GLOBE_R2023A_54009_{GHS_RES[res]}_V1_0"


def ghs_tile_url(epoch: int, res: int) -> str:
    n = ghs_name(epoch, res)
    return f"{GHS_BASE}/GHS_POP_E{epoch}_GLOBE_R2023A_54009_{GHS_RES[res]}/V1-0/tiles/{n}_{GHS_TILE}.zip"


def ghs_s3_url(epoch: int, res: int) -> str:
    return (f"https://jrc-ghsl.s3.amazonaws.com/ghs-pop/r2023a/54009/{res}m/{epoch}/"
            f"{ghs_name(epoch, res)}.tif")


GHS_PRODUCTS = [(2020, 100), (2025, 100), (2030, 100), (2025, 1000), (2030, 1000)]
GHS_DOCS = {
    "copyright.txt": f"{GHS_BASE}/GHS_POP_E2025_GLOBE_R2023A_54009_100/V1-0/copyright.txt",
    "GHSL_Data_Package_2023.pdf": "https://human-settlement.emergency.copernicus.eu/documents/GHSL_Data_Package_2023.pdf",
    "jrc_katalog_ghs_pop_r2023a.html": "https://data.jrc.ec.europa.eu/dataset/2ff68a52-5b5b-4a22-8f40-c41da8332cfe",
}
GHS_NODATA = -200.0

WP_URL = ("https://data.worldpop.org/GIS/Population/Global_2015_2030/R2025A/{y}/TUR/v1/100m/constrained/"
          "tur_pop_{y}_CN_100m_R2025A_v1.tif")
WP_YEARS = [2025, 2030]
WP_DOC_URL = "https://data.worldpop.org/repo/prj/Global_2015_2030/R2025A/doc/Global2_Release_Statement_R2025A_v1.pdf"
WP_CENSUS_URL = ("https://data.worldpop.org/repo/prj/Global_2015_2030/R2025A/doc/census/"
                 "global2_census_data_sources_R2025A_v1.xlsx")
WP_DOI = "10.5258/SOTON/WP00839"   # arama motorundan; açılış sayfası (hub.worldpop.org) vekil sunucuda engelli

KONTUR_URL = "https://geodata-eu-central-1-kontur-public.s3.amazonaws.com/kontur_datasets/kontur_population_TR_20231101.gpkg.gz"
KONTUR_GPKG = KONTUR_DIR / "kontur_population_TR_20231101.gpkg"


# ---------------------------------------------------------------- indirme yardımcıları
def _manifest_record(dest: Path, url: str) -> dict | None:
    man = dest.parent / "_manifest.jsonl"
    if not (dest.exists() and man.exists()):
        return None
    rel = str(dest.relative_to(common.ROOT))
    for line in man.read_text(encoding="utf-8").splitlines():
        try:
            rec = json.loads(line)
        except json.JSONDecodeError:
            continue
        if rec.get("url") == url and rec.get("file") == rel:
            return rec
    return None


def ensure(url: str, dest: Path, *, timeout: int = 600, overwrite: bool = False) -> dict:
    """Dosya yoksa indirir; manifestte (url, dosya) kaydı varsa tekrar yazmaz."""
    rec = None if overwrite else _manifest_record(dest, url)
    if rec is None:
        rec = common.download(url, dest, timeout=timeout, overwrite=overwrite)
        print(f"  kayıt: {rec['file']} ({rec['bytes']:,} B)")
    return rec


PHASE0_JSON = WP_DIR / "_faz0_kontrol.json"


def ensure_verified_existing(url: str, dest: Path) -> dict:
    """Faz 0'dan kalan dosyayı kaynaktan yeniden indirip sha256 karşılaştırır (yalnız ilk çalıştırmada).

    Sonuç PHASE0_JSON'da saklanır; sonraki çalıştırmalar oradan okur.
    """
    checks = json.loads(PHASE0_JSON.read_text(encoding="utf-8")) if PHASE0_JSON.exists() else {}
    rel = str(dest.relative_to(common.ROOT))
    rec = _manifest_record(dest, url)
    if rec is None:
        old = common.sha256(dest) if dest.exists() else None
        rec = common.download(url, dest, overwrite=True)
        checks[rel] = {"phase0_sha256": old, "downloaded_sha256": rec["sha256"],
                       "phase0_sha256_match": (old == rec["sha256"]) if old else None,
                       "checked": common.today()}
        PHASE0_JSON.write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"  yeniden indirildi: {rec['file']} — Faz 0 kopyasıyla aynı: {checks[rel]['phase0_sha256_match']}")
    return {**rec, "phase0_sha256_match": checks.get(rel, {}).get("phase0_sha256_match")}


def check_gunzip(gz: Path, out: Path) -> dict:
    """gz içeriğinin sha256'sını akışta hesaplar; açılmış dosya yoksa yazar, varsa karşılaştırır."""
    h = hashlib.sha256()
    tmp = out.with_suffix(out.suffix + ".part")
    write = not out.exists()
    with gzip.open(gz, "rb") as fi:
        fo = open(tmp, "wb") if write else None
        for block in iter(lambda: fi.read(1 << 20), b""):
            h.update(block)
            if fo:
                fo.write(block)
        if fo:
            fo.close()
            tmp.replace(out)
    h_out = common.sha256(out)
    return {"file": str(out.relative_to(common.ROOT)), "sha256": h_out, "matches_gz_content": h_out == h.hexdigest()}


# ---------------------------------------------------------------- GHS-POP
def _ghs_zip(epoch: int, res: int) -> Path:
    return GHS_DIR / f"{ghs_name(epoch, res)}_{GHS_TILE}.zip"


def _int_window(bounds, transform, expect: tuple[int, int]):
    """Sınırları tam sayılı piksel penceresine çevirir; hizasızlıkta hata verir."""
    from rasterio.windows import Window, from_bounds
    w = from_bounds(*bounds, transform=transform)
    vals = [w.col_off, w.row_off, w.width, w.height]
    if any(abs(v - round(v)) > 1e-6 for v in vals):
        raise ValueError(f"pencere ızgaraya hizalı değil: {w}")
    w = Window(*(int(round(v)) for v in vals))
    if (w.height, w.width) != expect:
        raise ValueError(f"pencere boyutu {w.height}x{w.width}, beklenen {expect}")
    return w


def read_ghs_window(path: str, res: int):
    """GHS rasterından Ankara penceresini (ham, float64, nodata=-200 korunur) okur."""
    import rasterio
    xmin, ymin, xmax, ymax = common.GHS_WINDOW
    expect = (int((ymax - ymin) / res), int((xmax - xmin) / res))
    with rasterio.open(path) as src:
        assert abs(src.res[0] - res) < 1e-9 and abs(src.res[1] - res) < 1e-9, src.res
        assert "Mollweide" in src.crs.to_wkt(), src.crs
        assert src.nodata == GHS_NODATA, src.nodata
        w = _int_window(common.GHS_WINDOW, src.transform, expect)
        a = src.read(1, window=w)
        return a, src.window_transform(w), src.crs, src.dtypes[0]


def write_tif(path: Path, arr: np.ndarray, transform, crs, tags: dict) -> None:
    import rasterio
    path.parent.mkdir(parents=True, exist_ok=True)
    prof = {"driver": "GTiff", "height": arr.shape[0], "width": arr.shape[1], "count": 1,
            "dtype": "float32", "crs": crs, "transform": transform, "compress": "deflate",
            "predictor": 3, "zlevel": 9, "tiled": True, "blockxsize": 256, "blockysize": 256}
    with rasterio.open(path, "w", **prof) as dst:
        dst.write(arr.astype(np.float32), 1)
        dst.update_tags(**{k: str(v) for k, v in tags.items()})


def ghs_products() -> dict:
    out = {}
    for e, r in GHS_PRODUCTS:
        zp = _ghs_zip(e, r)
        a, tr, crs, dt = read_ghs_window(f"/vsizip/{zp}/{ghs_name(e, r)}_{GHS_TILE}.tif", r)
        n_nodata = int((a == GHS_NODATA).sum())
        clean = np.where(a == GHS_NODATA, 0.0, a)
        tag = "100m" if r == 100 else "1km"
        dest = POP_DIR / f"ghs_pop_e{e}_{tag}.tif"
        write_tif(dest, clean, tr, crs, {
            "kaynak": f"{ghs_name(e, r)} (JRC GHS-POP R2023A), karo {GHS_TILE}",
            "url": ghs_tile_url(e, r), "birim": "kişi / hücre",
            "not": "Ankara penceresi (common.GHS_WINDOW); nodata -200 -> 0; float64 -> float32"
                   + ("; 2025 ve 2030 projeksiyondur" if e >= 2025 else ""),
            "atif": "Schiavina M., Freire S., Carioli A., MacManus K. (2023) doi:10.2905/2FF68A52-5B5B-4A22-8F40-C41DA8332CFE",
            "lisans": "CC BY 4.0"})
        out[(e, r)] = {"arr": clean, "raw": a, "transform": tr, "crs": crs, "file": dest,
                       "source_dtype": dt, "nodata_px": n_nodata,
                       "float32_total_diff": float(clean.astype(np.float32).sum(dtype=np.float64) - clean.sum())}
        print(f"  {dest.name}: {clean.shape[1]}x{clean.shape[0]} px, pencere toplamı {clean.sum():,.0f}, "
              f"nodata {n_nodata:,} px, kaynak tipi {dt}")
    return out


def check_s3_copy(ghs: dict) -> dict:
    """AWS kopyasını (jrc-ghsl.s3) resmî karo ile aynı pencerede piksel piksel karşılaştırır."""
    import rasterio
    import requests
    if S3_CHECK_JSON.exists():
        prev = json.loads(S3_CHECK_JSON.read_text(encoding="utf-8"))
        if len(prev.get("checks", [])) == len(GHS_PRODUCTS):
            return prev
    checks = []
    env = {"GDAL_DISABLE_READDIR_ON_OPEN": "EMPTY_DIR", "CPL_VSIL_CURL_ALLOWED_EXTENSIONS": ".tif",
           "GDAL_HTTP_MAX_RETRY": "3", "GDAL_HTTP_RETRY_DELAY": "2"}
    for e, r in GHS_PRODUCTS:
        url = ghs_s3_url(e, r)
        h = requests.head(url, headers={"User-Agent": common.USER_AGENT}, timeout=60)
        rec = {"product": ghs_name(e, r), "url": url, "accessed": common.today(), "http": h.status_code,
               "etag": h.headers.get("ETag"), "last_modified": h.headers.get("Last-Modified"),
               "content_length": h.headers.get("Content-Length"), "window": list(common.GHS_WINDOW)}
        if h.status_code == 200:
            with rasterio.Env(**env):
                a, tr, _, dt = read_ghs_window(f"/vsicurl/{url}", r)
            ref = ghs[(e, r)]["raw"]
            rec.update({"dtype": dt, "shape": list(a.shape),
                        "transform_equal": tuple(tr) == tuple(ghs[(e, r)]["transform"]),
                        "pixel_equal": bool(np.array_equal(a, ref)),
                        "max_abs_diff": float(np.nanmax(np.abs(a - ref))),
                        "window_sha256_s3": hashlib.sha256(a.tobytes()).hexdigest(),
                        "window_sha256_official": hashlib.sha256(ref.tobytes()).hexdigest()})
        checks.append(rec)
        print(f"  S3 {ghs_name(e, r)}: piksel eşit = {rec.get('pixel_equal')}")
    res = {"note": "Resmî JRC karosu (jeodpp) ile JRC AWS kopyası aynı Ankara penceresinde karşılaştırıldı.",
           "checks": checks}
    S3_CHECK_JSON.write_text(json.dumps(res, ensure_ascii=False, indent=2), encoding="utf-8")
    return res


def ghs_input_metadata_tur() -> dict:
    """GHS-POP R2023A girdi meta verisinden (zip içindeki xlsx) Türkiye satırı."""
    import io
    import openpyxl
    with zipfile.ZipFile(_ghs_zip(2025, 100)) as z:
        wb = openpyxl.load_workbook(io.BytesIO(z.read("GHS_POP_GLOBE_R2023A_input_metadata.xlsx")),
                                    read_only=True)
    rows = list(wb["INPUT METADATA"].iter_rows(values_only=True))
    hdr = rows[2]
    tur = next(r for r in rows if len(r) > 2 and r[2] == "TUR")
    names = []
    for i, h in enumerate(hdr):
        names.append(h if h not in names else f"{h}_{i}")
    return {k: v for k, v in zip(names, tur) if v not in (None, "")}


# ---------------------------------------------------------------- WorldPop
def clip_worldpop(year: int) -> dict:
    import rasterio
    from rasterio.windows import Window
    src_path = WP_DIR / f"tur_pop_{year}_CN_100m_R2025A_v1.tif"
    lon0, lat0, lon1, lat1 = common.ANKARA_BBOX
    with rasterio.open(src_path) as src:
        inv = ~src.transform
        c0, r0 = inv * (lon0, lat1)
        c1, r1 = inv * (lon1, lat0)
        c0, r0 = int(np.floor(c0 + 1e-6)), int(np.floor(r0 + 1e-6))
        c1, r1 = int(np.ceil(c1 - 1e-6)), int(np.ceil(r1 - 1e-6))
        w = Window(c0, r0, c1 - c0, r1 - r0)
        a = src.read(1, window=w).astype(np.float64)
        tr = src.window_transform(w)
        nod = src.nodata
        crs = src.crs
        tags = src.tags()
    n_nodata = int((a == nod).sum())
    clean = np.where(a == nod, 0.0, a)
    dest = POP_DIR / f"worldpop_{year}_100m.tif"
    write_tif(dest, clean, tr, crs, {
        "kaynak": f"WorldPop R2025A v1, tur_pop_{year}_CN_100m (constrained)", "url": WP_URL.format(y=year),
        "birim": "kişi / hücre", "not": f"yerel WGS84 3\" ızgara; common.ANKARA_BBOX kutusu; nodata {nod} -> 0",
        "atif": "Bondarenko M. ve ark. (2025) WorldPop R2025A v1, doi:10.5258/SOTON/WP00839",
        "lisans": tags.get("TIFFTAG_COPYRIGHT", "")})
    print(f"  {dest.name}: {clean.shape[1]}x{clean.shape[0]} px, kutu toplamı {clean.sum():,.0f}")
    return {"arr": clean, "transform": tr, "crs": crs, "file": dest, "nodata_px": n_nodata, "tags": tags}


def worldpop_census_tur() -> dict:
    """WorldPop R2025A nüfus girdi tablosundan Türkiye satırı (iki zaman noktası)."""
    import openpyxl
    wb = openpyxl.load_workbook(WP_DIR / "global2_census_data_sources_R2025A_v1.xlsx", read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    hdr = rows[2]
    tur = next(r for r in rows if len(r) > 2 and r[2] == "TUR")
    out = {}
    for i, (h, v) in enumerate(zip(hdr, tur)):
        if h is None or v in (None, ""):
            continue
        grp = "t1_" if 4 <= i <= 12 else "t2_" if 13 <= i <= 21 else ""
        out[f"{grp}{h.strip()}"] = v
    return out


def doi_resolves(doi: str) -> dict:
    """DOI çözümleyicisinin yönlendirdiği adresi kaydeder (açılış sayfası okunamıyorsa en azından kayıtlı mı)."""
    import requests
    try:
        r = requests.head(f"https://doi.org/{doi}", allow_redirects=False, timeout=30)
        return {"doi": doi, "http": r.status_code, "location": r.headers.get("Location"), "checked": common.today()}
    except requests.RequestException as exc:  # ağ yoksa
        return {"doi": doi, "error": str(exc)}


# ---------------------------------------------------------------- Kontur
def load_kontur_points():
    import geopandas as gpd
    import h3
    import pyogrio
    from pyproj import Transformer
    t = Transformer.from_crs(common.CRS_WGS84, "EPSG:3857", always_xy=True)
    lon0, lat0, lon1, lat1 = common.ANKARA_BBOX
    x0, y0 = t.transform(lon0, lat0)
    x1, y1 = t.transform(lon1, lat1)
    df = pyogrio.read_dataframe(KONTUR_GPKG, columns=["h3", "population"], bbox=(x0, y0, x1, y1),
                                read_geometry=True)
    res = sorted({h3.get_resolution(h) for h in df["h3"]})
    ll = np.array([h3.cell_to_latlng(h) for h in df["h3"]])
    pts = gpd.GeoDataFrame({"h3": df["h3"].values, "population": df["population"].values},
                           geometry=gpd.points_from_xy(ll[:, 1], ll[:, 0]), crs=common.CRS_WGS84)
    info = {"mean_cell_area_km2_bbox": float(np.mean([h3.cell_area(h, unit="km^2") for h in df["h3"]])),
            "h3_avg_area_km2_r8": float(h3.average_hexagon_area(8, unit="km^2")),
            "h3_avg_edge_km_r8": float(h3.average_hexagon_edge_length(8, unit="km"))}
    return pts, res, info


# ---------------------------------------------------------------- ilçe toplamları
def zonal_raster(arr: np.ndarray, transform, crs, districts, province) -> dict:
    """Piksel merkezi poligonda kuralıyla ilçe toplamları (1..n etiket) ve il/ilçe dışı kalan."""
    from rasterio.features import rasterize
    d = districts.to_crs(crs)
    p = province.to_crs(crs)
    lab = rasterize([(g, i + 1) for i, g in enumerate(d.geometry)], out_shape=arr.shape,
                    transform=transform, fill=0, dtype="int16", all_touched=False)
    inil = rasterize([(p.geometry.iloc[0], 1)], out_shape=arr.shape, transform=transform,
                     fill=0, dtype="uint8", all_touched=False).astype(bool)
    sums = np.bincount(lab.ravel(), weights=arr.ravel(), minlength=len(d) + 1)
    return {"district": sums[1:], "province": float(arr[inil].sum()),
            "in_province_no_district": float(arr[inil & (lab == 0)].sum()),
            "in_district_outside_province": float(arr[~inil & (lab > 0)].sum()),
            "window_total": float(arr.sum())}


def zonal_points(pts, districts, province) -> dict:
    import geopandas as gpd
    j = gpd.sjoin(pts, districts[["ilce_key", "geometry"]].reset_index(), how="left", predicate="within")
    j = j[~j.index.duplicated(keep="first")]
    s = j.groupby("index")["population"].sum()
    dist = np.array([float(s.get(i, 0.0)) for i in range(len(districts))])
    inp = pts.within(province.geometry.iloc[0])
    nod = j["index"].isna()
    return {"district": dist, "province": float(pts.loc[inp, "population"].sum()),
            "in_province_no_district": float(pts.loc[inp & nod.values, "population"].sum()),
            "in_district_outside_province": float(pts.loc[~inp & ~nod.values, "population"].sum()),
            "window_total": float(pts["population"].sum())}


def main() -> None:
    import geopandas as gpd

    print("1) İndirmeler")
    for e, r in GHS_PRODUCTS:
        ensure(ghs_tile_url(e, r), _ghs_zip(e, r))
    for fname, url in GHS_DOCS.items():
        ensure(url, GHS_DIR / fname)
    wp_recs = {}
    for y in WP_YEARS:
        dest = WP_DIR / f"tur_pop_{y}_CN_100m_R2025A_v1.tif"
        wp_recs[y] = ensure_verified_existing(WP_URL.format(y=y), dest)
    ensure(WP_DOC_URL, WP_DIR / "Global2_Release_Statement_R2025A_v1.pdf")
    ensure(WP_CENSUS_URL, WP_DIR / "global2_census_data_sources_R2025A_v1.xlsx")
    kgz = ensure(KONTUR_URL, KONTUR_DIR / (KONTUR_GPKG.name + ".gz"))
    kchk = check_gunzip(KONTUR_DIR / (KONTUR_GPKG.name + ".gz"), KONTUR_GPKG)
    print("  Kontur gpkg = gz içeriği:", kchk["matches_gz_content"])

    print("2) GHS-POP kırpma")
    ghs = ghs_products()
    s3 = check_s3_copy(ghs)
    tur_meta = ghs_input_metadata_tur()

    print("3) WorldPop kırpma")
    wp = {y: clip_worldpop(y) for y in WP_YEARS}

    print("4) Kontur")
    kpts, kres, kinfo = load_kontur_points()
    print(f"  {len(kpts):,} altıgen; H3 çözünürlüğü {kres}; ortalama hücre alanı (kutu) "
          f"{kinfo['mean_cell_area_km2_bbox']:.3f} km²; H3 r8 küresel ort. {kinfo['h3_avg_area_km2_r8']:.4f} km², "
          f"kenar {kinfo['h3_avg_edge_km_r8']:.3f} km")

    print("5) İlçe toplamları")
    d = gpd.read_file(common.PROCESSED / "ankara_ilceler.geojson").sort_values("ilce_key").reset_index(drop=True)
    il = gpd.read_file(common.PROCESSED / "ankara_il.geojson")
    assert len(d) == 25
    b = il.total_bounds
    lon0, lat0, lon1, lat1 = common.ANKARA_BBOX
    assert lon0 < b[0] and lat0 < b[1] and lon1 > b[2] and lat1 > b[3], "ANKARA_BBOX il sınırını kapsamıyor"
    gx0, gy0, gx1, gy1 = il.to_crs(common.CRS_MOLLWEIDE).total_bounds
    wx0, wy0, wx1, wy1 = common.GHS_WINDOW
    assert wx0 < gx0 and wy0 < gy0 and wx1 > gx1 and wy1 > gy1, "GHS_WINDOW il sınırını kapsamıyor"
    cols = {}
    for (e, r), g in ghs.items():
        cols[f"ghs_e{e}_{'100m' if r == 100 else '1km'}"] = zonal_raster(g["arr"], g["transform"], g["crs"], d, il)
    for y, w in wp.items():
        cols[f"worldpop_{y}_100m"] = zonal_raster(w["arr"], w["transform"], w["crs"], d, il)
    cols["kontur_20231101_h3r8"] = zonal_points(kpts, d, il)

    rows = []
    for i, r in d.iterrows():
        rows.append({"satir_turu": "ilce", "ilce": r["ilce"], "ilce_key": r["ilce_key"], "area_km2": r["area_km2"],
                     **{c: v["district"][i] for c, v in cols.items()}})
    rows.append({"satir_turu": "ilce_toplami", "ilce": "25 ilçe toplamı", "ilce_key": "",
                 "area_km2": float(d["area_km2"].sum()), **{c: v["district"].sum() for c, v in cols.items()}})
    rows.append({"satir_turu": "il_ici_ilce_disi", "ilce": "İl poligonunda, hiçbir ilçede değil", "ilce_key": "",
                 "area_km2": np.nan, **{c: v["in_province_no_district"] for c, v in cols.items()}})
    rows.append({"satir_turu": "il_poligonu", "ilce": "Ankara il poligonu", "ilce_key": "ankara",
                 "area_km2": float(il["area_km2"].iloc[0]), **{c: v["province"] for c, v in cols.items()}})
    df = pd.DataFrame(rows)
    num = [c for c in cols]
    df[num] = df[num].round(1)
    df.to_csv(OUT_CSV, index=False, encoding="utf-8")
    with pd.option_context("display.width", 250, "display.max_columns", 20):
        print(df.drop(columns=["ilce_key"]).to_string(index=False))

    summary = {
        "generated": common.today(),
        "ghs": {f"E{e}_{r}m": {"file": str(g["file"].relative_to(common.ROOT)), "shape": list(g["arr"].shape),
                               "source_dtype": g["source_dtype"], "nodata_px": g["nodata_px"],
                               "window_total": float(g["arr"].sum()),
                               "float32_total_diff": g["float32_total_diff"]} for (e, r), g in ghs.items()},
        "ghs_1km_vs_100m_window_total": {str(e): float(ghs[(e, 1000)]["arr"].sum() - ghs[(e, 100)]["arr"].sum())
                                          for e in (2025, 2030)},
        "ghs_s3_check": [{k: c.get(k) for k in ("product", "pixel_equal", "transform_equal", "max_abs_diff")}
                         for c in s3["checks"]],
        "ghs_input_metadata_TUR": tur_meta,
        "worldpop": {str(y): {"file": str(w["file"].relative_to(common.ROOT)), "shape": list(w["arr"].shape),
                              "transform": list(w["transform"])[:6], "nodata_px": w["nodata_px"],
                              "tags": w["tags"], "download": {k: wp_recs[y].get(k) for k in
                                                              ("url", "sha256", "bytes", "phase0_sha256_match")}}
                     for y, w in wp.items()},
        "worldpop_census_TUR": worldpop_census_tur(),
        "worldpop_doi": doi_resolves(WP_DOI),
        "kontur": {"gz": kgz["file"], "gz_sha256": kgz["sha256"], **kchk, "h3_resolutions": kres,
                   **kinfo, "cells_in_bbox": int(len(kpts))},
        "totals": {c: {k: (float(v) if not isinstance(v, np.ndarray) else None) for k, v in t.items()
                       if k != "district"} | {"district_sum": float(t["district"].sum())} for c, t in cols.items()},
    }
    SUMMARY_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print("  özet:", SUMMARY_JSON.relative_to(common.ROOT))


if __name__ == "__main__":
    main()
